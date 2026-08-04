import io

import pytest
from docx import Document
from pypdf import PdfWriter

from legal_core.document_text import DocumentTextError, extract_text


def _docx_bytes(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    doc = Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    if table:
        t = doc.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, val in enumerate(row):
                t.cell(r, c).text = val
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_docx_paragraphs_and_table_extracted():
    data = _docx_bytes(["Madde 1 gizlilik", "Madde 2 tedbirler"],
                       table=[["Kategori", "Süre"], ["Kimlik", "10 yıl"]])
    text = extract_text(data, "docx")
    assert "Madde 1 gizlilik" in text
    assert "Madde 2 tedbirler" in text
    assert "Kimlik" in text and "10 yıl" in text


def test_empty_docx_returns_empty_string():
    assert extract_text(_docx_bytes([]), "docx").strip() == ""


def test_pdf_with_no_text_returns_empty_string():
    w = PdfWriter()
    w.add_blank_page(width=200, height=200)
    buf = io.BytesIO()
    w.write(buf)
    assert extract_text(buf.getvalue(), "pdf").strip() == ""


def test_corrupt_bytes_raise_document_text_error():
    with pytest.raises(DocumentTextError):
        extract_text(b"not a real docx", "docx")
