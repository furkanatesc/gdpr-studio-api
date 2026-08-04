"""Yüklenen belge (docx/pdf) → düz metin çıkarımı — saf, DB/HTTP bilmez.

python-docx ve pypdf FONKSİYON İÇİNDE lazy import edilir (legal_core saf çekirdek;
masaüstü paketlemesi bu opsiyonel bağımlılıklara zorlanmaz).
"""

from __future__ import annotations

import io
from typing import Literal


class DocumentTextError(Exception):
    """Belge açılamadı/çözümlenemedi (bozuk dosya, yanlış biçim)."""


def _extract_docx(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts: list[str] = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def extract_text(data: bytes, kind: Literal["docx", "pdf"]) -> str:
    """docx/pdf bytes → düz metin. Boş sonuç boş string döner (çağıran 422 verir)."""
    try:
        if kind == "docx":
            return _extract_docx(data)
        if kind == "pdf":
            return _extract_pdf(data)
    except DocumentTextError:
        raise
    except Exception as e:
        raise DocumentTextError(str(e)) from e
    raise DocumentTextError(f"Desteklenmeyen tür: {kind}")
